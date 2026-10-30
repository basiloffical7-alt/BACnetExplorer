"""Login hashing / auto-login token and the .xlsx writer."""
import os, sys, tempfile, time, types
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
for n in ("tkinter", "tkinter.ttk"):                     # auth imports tkinter for the login form
    if n not in sys.modules:
        try: __import__(n)
        except ImportError:
            m = types.ModuleType(n); m.Frame = object; sys.modules[n] = m
sys.modules["tkinter"].ttk = sys.modules["tkinter.ttk"]
import authcore as auth
s = {}
a = auth.Auth(s)
assert a.check("admin", "Basil@12345") and a.check("ADMIN ", "Basil@12345")
assert not a.check("admin", "basil@12345") and not a.check("root", "Basil@12345")
assert "Basil@12345" not in str(s), "password must not be stored in clear text"
assert not a.autologin_ok(); a.enable_autologin(); assert a.autologin_ok()
assert a.change_password("wrong", "NewPass123") is not None
assert a.change_password("Basil@12345", "short") is not None
assert a.change_password("Basil@12345", "NewPass123") is None
assert not a.autologin_ok(), "auto login must be cancelled by a password change"
assert a.check("admin", "NewPass123") and not a.check("admin", "Basil@12345")
a.enable_autologin(); orig = auth.machine_key
auth.machine_key = lambda: "other-pc|user"; assert not a.autologin_ok(), "token is bound to the PC"
auth.machine_key = orig; assert a.autologin_ok()
tok = a.make_restart_token(); assert a.use_restart_token(tok) and not a.use_restart_token(tok)
tok = a.make_restart_token(); assert not a.use_restart_token("forged")
print("auth: ok")

from xlsx_export import Sheet, write_xlsx
path = os.path.join(tempfile.mkdtemp(), "r.xlsx")
sh = Sheet("Devices", ["Instance", "Name", "Status"])
sh.add([1001, "AHU <1> & 'x'", "Online"], {2: "green"}); sh.add([1002, "VAV", "Offline"], {2: "red"})
write_xlsx(path, [Sheet("Summary", ["Item", "Value"]), sh], title="t")
try:
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb["Devices"]
    assert ws["A2"].value == 1001 and ws["B2"].value == "AHU <1> & 'x'" and ws["C3"].value == "Offline"
    assert ws["C2"].fill.fgColor.rgb.endswith("C8E6C9") and ws["A1"].font.b and ws.freeze_panes == "A2"
    print("xlsx: ok (verified with openpyxl)")
except ImportError:
    print("xlsx: written (openpyxl not installed, not verified)")
