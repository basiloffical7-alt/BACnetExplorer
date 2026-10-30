"""Android app logic (android/core.py + graphics_core.py) against the simulator - real UDP."""
import os, sys, subprocess, tempfile, threading, time
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [os.path.join(ROOT, "android"), ROOT]
from core import Core, parse_value, value_text
from graphics_core import GraphicsModel
from bacnet.codec import ObjectId, Real
from bacnet import codec as X

sim = subprocess.Popen([sys.executable, "bacnet_simulator.py", "--port", "47809"], cwd=ROOT,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1)
results = {}
def wait(key, t=15):
    end = time.time() + t
    while key not in results and time.time() < end: time.sleep(0.05)
    assert key in results, f"timeout waiting for {key}"
    return results.pop(key)
data = tempfile.mkdtemp()
try:
    core = Core(data)                                     # post() runs inline in tests
    assert core.auth.check("admin", "Basil@12345")
    err = core.connect("127.0.0.1", "127.255.255.255", port=47871, remote_port=47809)
    assert err is None, err
    core.discover(("127.0.0.1", 47809))
    end = time.time() + 6
    while len(core.devices()) < 9 and time.time() < end: time.sleep(0.1)
    print("devices:", [(d.instance, d.name) for d in core.devices()])
    assert len(core.devices()) == 9
    dev = core.device(3001)
    prog = []
    core.load_objects(dev, prog.append, lambda r: results.__setitem__("objs", r), lambda e: results.__setitem__("objs", e))
    objs = wait("objs")
    assert not isinstance(objs, Exception), objs
    print(f"objects of {dev.name}: {len(objs)}  e.g.", [(p.oid.short, p.name, p.text()) for p in objs[:4]], "progress msgs:", len(prog))
    assert all(p.name for p in objs) and any(p.text().endswith("%") for p in objs)
    # read point
    core.read_point(dev, ObjectId(1, 2), lambda r: results.__setitem__("pt", r), lambda e: results.__setitem__("pt", e))
    props = wait("pt"); print("point props:", len(props), [X.format_value(r.value, 1, r.prop) for r in props if r.prop == 85])
    # parse + write analog / binary / release
    assert parse_value(ObjectId(1, 2), "65,5 %") == Real(65.5)
    assert int(parse_value(ObjectId(4, 1), "inactive")) == 0 and isinstance(parse_value(ObjectId(1, 1), "release"), X.Null)
    core.write(dev, ObjectId(1, 2), Real(61.0), 8, lambda v: results.__setitem__("w", v), lambda e: results.__setitem__("w", e))
    print("write analog -> read back:", wait("w"))
    core.write(dev, ObjectId(4, 1), X.Enumerated(0), 8, lambda v: results.__setitem__("w", v), lambda e: results.__setitem__("w", e))
    v = wait("w"); print("write binary -> read back:", value_text(v, ObjectId(4, 1))); assert int(v) == 0
    # watch list + flash
    core.add_watch(dev, ObjectId(0, 1), "Supply Air Temp"); core.add_watch(dev, ObjectId(4, 1), "Supply Fan Command")
    core.add_watch(core.device(3002), ObjectId(19, 1), "Fan Speed")
    time.sleep(3)
    for w in core.watch: print(f"   watch {w.dev} {w.oid.short:<5} {w.name:<20} {w.text:<12} {w.state}")
    assert all(w.text != "-" for w in core.watch)
    core.write(dev, ObjectId(4, 1), X.Enumerated(1), 8, lambda v: results.__setitem__("w", v), lambda e: results.__setitem__("w", e)); wait("w")
    fan = next(w for w in core.watch if w.oid == ObjectId(4, 1))
    print("after writing fan ON: text", fan.text, "flashing", fan.flashing); assert fan.flashing and fan.text == "active"
    # persistence
    core.save(); c2 = Core(data); assert [w.name for w in c2.watch] == [w.name for w in core.watch]
    # online report + excel
    cancel = threading.Event()
    core.scan(True, False, lambda *a: None, lambda r: results.__setitem__("rep", r), cancel)
    rep = wait("rep", 40); cnt = core.report_counts(rep); print("report:", cnt)
    assert cnt["dev_on"] == 9 and cnt["pt_total"] > 250
    xl = os.path.join(data, "report.xlsx"); core.export_report(rep, xl)
    import openpyxl; wb = openpyxl.load_workbook(xl); print("xlsx sheets:", wb.sheetnames, "points rows:", wb["Points"].max_row - 1)
    # graphics
    g = GraphicsModel(core)
    g.add_page("CH plant", "chiller"); g.set_options({"chillers": 2, "primary": 2, "secondary": 2})
    sugg = g.suggest_links(core.device(3005).object_names or {})
    if not sugg:
        core.load_objects(core.device(3005), lambda m: None, lambda r: results.__setitem__("o5", r), lambda e: results.__setitem__("o5", e)); wait("o5")
        sugg = g.suggest_links(core.device(3005).object_names)
    g.apply_links(3005, sugg); print("chiller auto-linked:", len(sugg), "of", len(g.template().POINTS))
    assert len(sugg) == len(g.template().POINTS)
    time.sleep(3); g.sync(); time.sleep(2)
    rows = g.point_rows(); print("graphic rows e.g.:", [(r["label"], r["value"]) for r in rows[:5]])
    assert all(r["value"] not in ("", "waiting…") for r in rows)
    rec, d = g.render(1000, 520, dark=False, demo=False, t=3.0)
    print("graphic drawn:", len(rec.ops), "items; tap at chiller 1 ->", rec.point_at(d.X(g.template().ch_x(0)), d.Y(265)))
    sys.path.insert(0, os.path.join(ROOT, "tests")); from render_canvas import RecCanvas
    rc = RecCanvas(w=1000, h=520); rc.ops = rec.ops; rc.render(os.environ.get("PNG", "/tmp/android_graphic.png"))
    core.close(); c2.close()
    print("ANDROID CORE TESTS PASSED")
finally:
    sim.terminate()
